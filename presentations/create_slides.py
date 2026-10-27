from pathlib import Path
import math
from PIL import ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement

OUT = Path(__file__).resolve().parent
prs = Presentation()
prs.slide_width = Inches(13.333333)
prs.slide_height = Inches(7.5)
BG='F6F3EC'; INK='183244'; MUTED='61727D'; TEAL='007F78'; GOLD='E4AB38'
CODE='132F3D'; WHITE='FFFFFF'; PALE='E4EFEB'; RULE='D8DFDA'; RED='B45747'
FONT='Noto Sans CJK JP'; MONO='DejaVu Sans Mono'
JP_PATH='/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc'
MONO_PATH='/usr/share/fonts/truetype/dejavu/DejaVuSansMono.ttf'
notes=[]
checks=[]

def rgb(value): return RGBColor.from_string(value)

def rect(s,x,y,w,h,fill,outline=None,radius=False):
    sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(fill)
    if outline:
        sh.line.color.rgb=rgb(outline)
    else: sh.line.fill.background()
    if radius:
        try: sh.adjustments[0]=0.12
        except Exception: pass
    return sh

def text(s,x,y,w,h,value,size=24,color=INK,bold=False,font=FONT,align=None,fit=True):
    sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h))
    tf=sh.text_frame; tf.clear(); tf.word_wrap=False
    tf.margin_left=tf.margin_right=Inches(.015)
    tf.margin_top=tf.margin_bottom=Inches(.015)
    lines=value.split('\n')
    actual=size
    if fit:
        path=MONO_PATH if font==MONO else JP_PATH
        f=ImageFont.truetype(path, size=round(size*4/3))
        maximum=max((f.getlength(line) for line in lines),default=0)
        if maximum > (w-.03)*96*.93:
            actual=math.floor(size*(w-.03)*96*.93/maximum*10)/10
    for i,line in enumerate(lines):
        p=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        p.text=line; p.space_after=Pt(4); p.space_before=Pt(0)
        p.line_spacing=1.16
        if align is not None: p.alignment=align
        for r in p.runs:
            r.font.name=font; r.font.size=Pt(actual); r.font.bold=bold; r.font.color.rgb=rgb(color)
            rp=r._r.get_or_add_rPr(); rp.set('lang','ja-JP')
            ea=OxmlElement('a:ea'); ea.set('typeface',FONT); rp.append(ea)
    checks.append((len(prs.slides),value[:35],actual))
    return sh

def line(s,x1,y1,x2,y2,color=TEAL,width=2):
    sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,Inches(x1),Inches(y1),Inches(x2),Inches(y2))
    sh.line.color.rgb=rgb(color); sh.line.width=Pt(width)
    return sh

def arrow(s,x,y,w=.45,h=.28,color=TEAL):
    return rect_shape(s,MSO_SHAPE.RIGHT_ARROW,x,y,w,h,color)

def rect_shape(s,kind,x,y,w,h,color):
    sh=s.shapes.add_shape(kind,Inches(x),Inches(y),Inches(w),Inches(h))
    sh.fill.solid(); sh.fill.fore_color.rgb=rgb(color); sh.line.fill.background(); return sh

def label(s,x,y,w,head,body='',h=1.35,fill=WHITE,accent=TEAL):
    rect(s,x,y,w,h,fill,radius=True)
    rect(s,x,y,.06,h,accent)
    compact=bool(body) and h<=1.05
    text(s,x+.2,y+(.1 if compact else .14),w-.38,.42,head,21 if compact else 23,accent,True)
    if body: text(s,x+.2,y+(.53 if compact else .68),w-.38,h-(.57 if compact else .72),body,18 if compact else 20)

def code(s,x,y,w,lines,size=23,title=None):
    if isinstance(lines,str): lines=lines.split('\n')
    header=.47 if title else .12
    h=header+.18+len(lines)*(size/72*1.23+.025)
    rect(s,x,y,w,h,CODE,radius=True)
    if title: text(s,x+.22,y+.12,w-.44,.28,title,12,'8FBCB7',True)
    for i,value in enumerate(lines):
        text(s,x+.23,y+header+.08+i*(size/72*1.23+.025),w-.46,size/72*1.24,
             value,size,'EAF3F1',font=MONO)
    return h

def base(tag,title,subtitle='',source=''):
    s=prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb=rgb(BG)
    rect(s,.45,.35,.12,.32,TEAL)
    text(s,.72,.33,11.8,.32,tag,12,TEAL,True)
    text(s,.72,.96,11.9,.75,title,33,INK,True)
    if subtitle: text(s,.75,1.83,11.8,.5,subtitle,18,MUTED)
    line(s,.72,6.98,12.62,6.98,RULE,1)
    text(s,.75,7.08,11,.2,source or 'Arduino Uno / ATmega328P  |  AES measurement → AVR assembly',9,MUTED)
    text(s,12,7.02,.6,.3,f'{len(prs.slides):02d}',12,TEAL,True,align=PP_ALIGN.RIGHT)
    return s

def note(s,body,source=''):
    content=body+ ('\n\n出典・確認箇所：'+source if source else '')
    s.notes_slide.notes_text_frame.text=content
    notes.append((len(prs.slides),s.shapes[2].text if len(s.shapes)>2 and s.shapes[2].has_text_frame else '',content))

# 01
s=base('BEGINNER GUIDE / 01–18 本編・19–21 補足','Arduino Uno の AES を\nアセンブリで読む','', '対象：今回作成した measurement.S と添付の Standard_AES_Measurement.ino')
s.shapes[2].height=Inches(1.45)
text(s,.78,2.8,10.9,.65,'C++ の「1行」を、CPU への「細かい指示」にする。',28,TEAL,True)
for x,num,head,body in [( .8,'01','知る','CPU・メモリの基礎'),(4.95,'02','比べる','どこを、どう変更したか'),(9.1,'03','読む','実際の命令を1行ずつ')]:
    label(s,x,4.25,3.45,head,body,h=1.45)
    text(s,x+.2,5.95,3,.35,num,14,TEAL,True)
note(s,'この資料は初学者が実際のアセンブリを読むための説明です。AESの安全性や暗号理論をすべて説明する資料ではありません。今回のコードは初期鍵とのXORと第1ラウンドの測定用です。元の測定機能を維持しながら、プログラムの処理を細かいCPU命令に移しました。目安は本編15〜20分、補足は必要なときに参照してください。')

# 02
s=base('STEP 1 / 元の動作を知る','まず、何を測っていたのか？','平文＝AES への入力データ。16バイトの平文を、同じ値で10回測定します。','原本：execute_aes_std()、loop() / measurement.S:285–292, 148–175')
stages=[('平文','16バイト'),('初期鍵XOR','Round 0'),('置き換え','SubBytes'),('並べ替え','ShiftRows'),('混合','MixColumns'),('鍵XOR','Round 1')]
for i,(head,body) in enumerate(stages):
    x=.8+i*2.06
    rect(s,x,2.95,1.74,1.24,WHITE,radius=True)
    text(s,x+.1,3.13,1.54,.4,head,20,TEAL,True,align=PP_ALIGN.CENTER)
    text(s,x+.08,3.7,1.58,.28,body,14,MUTED,align=PP_ALIGN.CENTER)
    if i<5: arrow(s,x+1.79,3.43,.22,.19)
line(s,2.84,4.65,12.78,4.65,TEAL,3)
text(s,2.85,4.86,9.7,.35,'D5 を HIGH にする → この計算区間を測定 → LOW にする',21,TEAL,True)
label(s,.8,5.62,11.88,'測定用の「初期 XOR ＋ 第1ラウンド」','AES-128 の全10ラウンド暗号化ではありません。',h=1.0,fill=PALE)
note(s,'平文は暗号処理に入れる元の16バイトのデータです。stateはその計算途中の状態を保持します。初期鍵とのXORの後にSubBytes、ShiftRows、MixColumns、第1ラウンド鍵とのXORを行います。D5のHIGHとLOWの間が測定対象です。総数は5000平文で、1回のSコマンドで500平文、各10回、計5000パルスを出します。対象を誤ってAES全体の暗号文と呼ばないようにします。','reference/Standard_AES_Measurement.ino:118–129, 172–267; measurement.S:148–175, 285–292')

# 03
s=base('STEP 2 / アセンブリ化の意味','両方とも、最後は「機械語」になる','アセンブリでは、CPU に行わせる小さな操作を人が明示します。')
for y,kind,middle,desc in [(2.9,'C++','コンパイラ','まとまった処理を書き、変換を任せる'),(4.5,'アセンブリ','アセンブラ','読む・計算する・保存する、を具体的に書く')]:
    label(s,.8,y,2.7,kind,h=1.0,fill=WHITE)
    arrow(s,3.72,y+.32,.6,.3)
    label(s,4.55,y,3.3,middle,h=1.0,fill=PALE)
    arrow(s,8.1,y+.32,.6,.3)
    label(s,8.95,y,3.52,'機械語',h=1.0,fill=WHITE)
    text(s,.97,y+1.09,11.4,.35,desc,20,MUTED)
note(s,'機械語はCPUが実際に実行する数値の命令です。元のArduinoのC++も既に機械語へコンパイルされて動いています。アセンブリ化は初めてCPUで動くようにすることではなく、処理を記述する抽象度を変えることです。GNU AVRアセンブリを使い、.Sファイルは前処理後にアセンブルされます。命令列を明示できることが目的で、速度向上を今回証明したわけではありません。')

# 04
s=base('STEP 2 / CPU の作業場所','レジスタ・SRAM・フラッシュを区別する','Uno は8ビットCPU。レジスタ1個には、通常1バイトを入れます。')
for x,head,body,eg in [( .8,'レジスタ','CPU 内の小さな作業場所','r18、r19 など'),(4.95,'SRAM','処理中に読み書きする場所','state、pt、round_keys'),(9.1,'フラッシュ','プログラムや定数を置く場所','命令、S-Box')]:
    label(s,x,2.92,3.43,head,body,h=1.75)
    text(s,x+.18,4.95,3.07,.48,eg,21,TEAL,True)
text(s,.98,5.66,11.5,.95,'state は、計算途中の16バイトのデータです。\nX・Z は「メモリのどこを見るか」を示す、アドレス用のレジスタ組です。',20)
note(s,'ビットは0か1の1桁、8ビットで1バイトです。ATmega328Pの汎用レジスタはr0からr31の32個で、それぞれ8ビットです。ここでは作業用のr18などに注目します。Xはr27:r26、Zはr31:r30の組で16ビットのアドレスを持ちます。XやZはデータそのものではなく住所を保持します。今回のS-Boxはフラッシュ、stateと鍵はSRAMにあります。','measurement.S:16–28, 302–307, 417–439')

# 05
s=base('STEP 3 / 1行の読み方','「命令 → 対象 → 説明」の順で読む','AVR の2オペランド命令では、一般に左側が結果の置き場所です。')
code(s,.88,2.85,11.55,['eor  r18, r19   ; XOR'],34,title='命令の読み方（コメントは説明用）')
for x,head,body in [( .9,'eor','XOR を行う命令'),(4.95,'r18, r19','r18 と r19 を使う'),(9.0,'; 以降','人が読むコメント')]:
    label(s,x,4.6,3.48,head,body,h=1.4)
text(s,.95,6.29,11.6,.3,'意味：r18 ← r18 XOR r19。r19 の値は変えず、結果を r18 に入れます。',21,TEAL,True)
note(s,'eorはexclusive ORの略で、AVRのXOR命令です。命令、空白、オペランド、コメントの順に読みます。r18とr19は変数名ではなくCPUの作業用レジスタ名です。左側が保存先という説明はeorやmovなどについての一般則で、st X+,r18のようなメモリ保存命令は個別に意味を確認します。セミコロン以降はアセンブラでは命令として実行されません。')

# 06
s=base('STEP 3 / XOR の意味','XOR は「同じなら0、違えば1」','2つの数を、ビットごとに比較します。足し算の繰り上がりはありません。')
for y,lead,bits,hx,c in [(2.85,'データ','0 0 1 1 1 1 0 0','0x3C',INK),(3.65,'鍵','1 0 1 0 0 1 0 1','0xA5',TEAL),(4.72,'結果','1 0 0 1 1 0 0 1','0x99',TEAL)]:
    text(s,1.0,y,1.35,.45,lead,24,c,True)
    text(s,2.7,y,6.5,.5,bits,30,c,True,font=MONO)
    text(s,10,y,2,.5,hx,27,c,True,font=MONO)
line(s,2.65,4.42,11.85,4.42,RULE,2)
text(s,.98,6.0,11.2,.45,'C++ の ^ と、AVR の eor は、どちらもこの XOR を表します。',24)
note(s,'例の0xは16進数を表す接頭辞です。3Cは00111100、A5は10100101なので、XORは10011001すなわち99になります。暗号の鍵とのXORはこの計算です。このスライドは命令の意味を理解するための単純な例で、実際の最初のデータ7Aと鍵2Bの例は次のスライドの後に示します。')

# 07
s=base('STEP 4 / 変更例① AddRoundKey','C++ の1行を、4つの命令にする','X は state の現在位置、Z は選んだラウンド鍵の現在位置を指す前提です。','原本:111 / measurement.S:309–312（ポインタ初期化・ループ部分を除いた抜粋）')
code(s,.88,2.77,11.55,['state[i] ^= round_keys[round][i];'],24,title='変更前 / C++')
rows=[('ld  r18, X','① state の1バイトを読む'),('ld  r19, Z+','② 鍵の1バイトを読む'),('eor r18, r19','③ XOR を計算する'),('st  X+, r18','④ 同じ state の位置へ書き戻す')]
for i,(inst,desc) in enumerate(rows):
    y=4.1+i*.57
    rect(s,.9,y,4.35,.48,CODE,radius=True)
    text(s,1.12,y+.05,3.9,.35,inst,22,'EAF3F1',font=MONO)
    text(s,5.7,y+.04,6.7,.4,desc,23,TEAL if i==2 else INK,bold=i==2)
note(s,'ここがアセンブリ化の最も分かりやすい例です。C++の1行にはメモリ読み込み、計算、メモリへの書き戻しが含まれます。それを明示しました。この抜粋の前でXとZのアドレスを設定し、外側のループで16バイト分処理します。Z+は鍵の読み込み後、X+はstateへの保存後にそれぞれアドレスを1進めます。この4行だけでAES全体が完了するわけではありません。','reference/Standard_AES_Measurement.ino:109–113; measurement.S:297–315')

# 08
s=base('STEP 4 / 住所とデータの動きを追う','X+ の「+」は、処理した後に住所を1進める','実際の最初の平文と初期鍵を、1バイトだけ追ってみます。')
label(s,.9,2.8,3.15,'SRAM のデータ','state[0] = 0x7A\n鍵[0]    = 0x2B',h=1.8)
arrow(s,4.22,3.5,.55,.3)
label(s,4.95,2.8,3.2,'CPU の作業場所','r18 = 0x7A\nr19 = 0x2B',h=1.8,fill=PALE)
arrow(s,8.34,3.5,.55,.3)
label(s,9.06,2.8,3.43,'XOR の結果','r18 = 0x51',h=1.8)
code(s,.94,5.0,4.25,['st X+, r18'],26)
text(s,5.6,5.02,6.9,1.1,'state[0] に 0x51 を保存\nその後、X は state[1] を指す',23)
text(s,.98,6.33,11.4,.32,'0x51 は初期 XOR の途中結果です。第1ラウンド終了後の先頭バイトは 0x1C です。',17,MUTED)
note(s,'ポインタは住所、レジスタの中身は読んできたデータ、と区別します。ld r18,XではXはまだ進みません。st X+,r18によって同じstate[0]を書き換えてから次へ進みます。鍵はld r19,Z+で読んだ時点で次の鍵へ進みます。7A XOR 2Bは51です。これはAESの途中結果であり、ログの最終値State_R1と混同しないようにします。','referenceの初回PTと鍵; measurement.S:309–312; build/asm-test/serial.log の Trace:0')

# 09
s=base('STEP 4 / 変更例② for ループ','繰り返しは「回数を減らし、戻る」で作る','AddRoundKey は、このループで state の16バイトを処理します。','measurement.S:308–315 / 1b は「直前にあるラベル1へ戻る」')
code(s,.9,2.8,5.45,['    ldi r20, 16','1:  ld  r18, X','    ld  r19, Z+','    eor r18, r19','    st  X+, r18','    dec r20','    brne 1b'],23,title='実際のコードから抜粋')
label(s,6.82,2.84,5.57,'ldi r20, 16','残り回数を16にする',h=1.0)
label(s,6.82,4.02,5.57,'dec r20','残り回数を1減らす',h=1.0)
label(s,6.82,5.2,5.57,'brne 1b','0でなければ、ラベル1に戻る',h=1.0,fill=PALE)
note(s,'C++のforは初期化、処理、カウンタ更新、条件判断をまとめて記述できます。ここでは残り回数方式です。ldiは数値を直接入れる命令、decは1減らす命令です。decで結果が0になるとゼロフラグが立ちます。brneはゼロフラグが立っていなければ分岐するため、この並びでは残り回数が0でないときに繰り返します。1bのbはbackwardで、数値1の直前のローカルラベルを探します。','measurement.S:308–315')

# 10
s=base('STEP 4 / 変更例③ SubBytes','S-Box は、フラッシュから1バイト読む','S-Box は「入力値に対応する値を引く、256個の対応表」です。','原本:59–62 / measurement.S:323–329')
label(s,.9,2.8,5.15,'変更前：pgm_read_byte(...)','C++ の関数形式で対応表を参照',h=1.35)
label(s,.9,4.42,5.15,'変更後：lpm r19, Z','Z の指すフラッシュを読む命令',h=1.35,fill=PALE)
code(s,6.45,2.8,5.96,['ld  r18, X','ldi r30, lo8(sbox_std)','ldi r31, hi8(sbox_std)','add r30, r18','adc r31, r1','lpm r19, Z','st  X+, r19'],20,title='実際のコード / 1バイト分')
text(s,.99,6.28,11.4,.34,'先頭アドレス ＋ 入力値 → 読む場所を決める → lpm で値を取り出す',20,TEAL,True)
note(s,'普通のSRAMを読むldと、フラッシュを読むlpmの使い分けに注目します。r30がZの下位8ビット、r31が上位8ビットです。lo8とhi8はアドレスをその2つに分けます。入力値を下位側へaddし、繰り上がりがあればadcで上位側に反映します。r1はAVR GCCの規約に従い0に保っています。表の中身は元の標準AES S-Boxと同じです。上級者向けのアドレス計算の詳細は発表時にすべて説明しなくても構いません。','measurement.S:319–332, 423–439')

# 11
s=base('STEP 4 / 変更例④ ShiftRows','並べ替えは「一時保存 → 読む → 書く」','例：行1の [state[1], state[5], state[9], state[13]] を左へ1つ回転。','原本:69–73 / measurement.S:337–344')
for y,head,vals in [(2.85,'変更前',['A','B','C','D']),(4.15,'変更後',['B','C','D','A'])]:
    text(s,.98,y+.16,1.7,.5,head,24,TEAL,True)
    for j,value in enumerate(vals):
        rect(s,3.1+j*1.3,y,1.04,.9,PALE if value!='A' else 'F8E8C4',radius=True)
        text(s,3.1+j*1.3,y+.16,1.04,.55,value,30,INK,True,align=PP_ALIGN.CENTER)
text(s,9.05,3.1,3.15,1.8,'最初の A を\nレジスタへ退避し、\n最後に書き戻す。',22)
code(s,.95,5.4,6.0,['lds r18, state+1','lds r19, state+5','sts state+1, r19'],19)
text(s,7.35,5.73,5.02,.83,'r18 が C++ の temp に相当。\n全行とも元と同じ並びになります。',20)
note(s,'A、B、C、Dはその場所に入っていたバイト値の説明用記号です。実際の添字の並びは1、5、9、13で、連続したSRAMの4バイトではありません。最初の値をr18へ保存してから次の値を前へ詰め、最後に退避した値をstate[13]へ置きます。他の行も元コードと同じ0、1、2、3個の回転です。ldsとstsは名前付きのSRAMアドレスを直接指定する読み書き命令です。','reference/Standard_AES_Measurement.ino:65–89; measurement.S:335–362')

# 12
s=base('STEP 4 / 変更例⑤ MixColumns','「左シフト ＋ 必要な XOR」で、列を混ぜる','AES の2倍演算は、8ビット内の左シフトと補正で計算できます。','原本:galois_mul2()、MixColumns_Std() / measurement.S:364–413')
code(s,.9,2.8,11.5,['(a << 1) ^ ((a & 0x80) ? 0x1B : 0x00)'],24,title='変更前 / galois_mul2 の中身')
code(s,.93,4.1,5.17,['mov  r24, r23','lsl  r23','sbrc r24, 7','eor  r23, r25'],23,title='XTIME r23,r24 の展開（r25 = 0x1B）')
text(s,6.57,4.35,5.67,1.9,'① 元の値を r24 に保存\n② r23 を左に1ビットずらす\n③ 元の最上位ビットが0なら次を飛ばす\n④ 1なら 0x1B と XOR して補正',20)
text(s,.98,6.54,11.6,.25,'列の式も同じ結果になる形へ整理。演算回数や波形が同じになる保証はありません。',16,MUTED)
note(s,'左シフトはビットを左にずらすことです。AESが使う有限体での2倍は通常の整数の2倍と異なり、元の最上位ビットが1だったときに1BとのXORで補正します。macro XTIMEのvalue=r23、scratch=r24とした実際の展開です。sbrcは指定ビットが0なら直後の1命令を飛ばします。MixColumnsではt=a0 XOR a1 XOR a2 XOR a3を一度求め、out[i]=a[i] XOR t XOR xtime(a[i] XOR a[next])という同値な式を使いました。これは式の整理も含む移植であり、逐語的な1対1変換だけではありません。','reference/Standard_AES_Measurement.ino:91–107; measurement.S:364–413')

# 13
s=base('STEP 4 / 変更例⑥ 32ビットの乱数','32ビットの値を、4つの8ビットに分ける','元の xorshift32 と同じ計算順とシードを維持しました。')
code(s,.9,2.83,5.52,['x ^= x << 13;','x ^= x >> 17;','x ^= x << 5;'],26,title='元の計算 / 初期シード12345')
for j,reg in enumerate(['r18','r19','r20','r21']):
    rect(s,6.95+j*1.32,3.22,1.08,.8,PALE,radius=True)
    text(s,6.96+j*1.32,3.44,1.06,.35,reg,20,TEAL,True,font=MONO,align=PP_ALIGN.CENTER)
text(s,7,4.32,5.05,.9,'下位 → 上位の順に4バイト。\n桁をまたぐ1ビットも引き継ぎます。',20)
code(s,.95,5.15,5.4,['lsl r22','rol r23','rol r24','rol r25'],17,title=None)
text(s,6.95,5.24,5.1,1.35,'複製した r22–r25 をシフト。\nlsl で出た1ビットを、\nrol が次のバイトへ渡します。',20)
note(s,'状態値はr18〜r21へロードします。シフトに使う複製はr22〜r25です。スライドの下の4命令はx<<13のループからの抜粋で、13回繰り返す本体です。lslで押し出されたビットはキャリーフラグに入り、次のrolがそれを最下位ビットへ取り込みます。右シフト側はlsrとrorで上位から下位へ行います。32ビットの状態全体を更新し、呼び出し側が使用する下位8ビットをr24で返します。戻り値全体を32ビットで返す関数へ変更したわけではありませんが、元のスケッチが使用する平文生成結果は同じです。','reference/Standard_AES_Measurement.ino:xorshift32(); measurement.S:230–281')

# 14
s=base('STEP 4 / 変更例⑦ if とバッチ制御','条件を比較し、次に進む場所を選ぶ','大文字 S で開始・再開。完了後の S ではシードも初期化します。','measurement.S:75–99, 177–212 / 比較の結果はフラグに記録される')
code(s,.92,2.85,5.25,["cpi  r24, 'S'",'breq 2f','ret','2:  ; start / resume'],23,title='受信済み文字を判定（最後のコメントは説明用）')
label(s,6.66,2.86,5.75,'cpi → 比較する','受信文字が S と同じか確認',h=1.16)
label(s,6.66,4.22,5.75,'breq → 等しければジャンプ','2f：前方にあるラベル2へ進む',h=1.16,fill=PALE)
for i,val in enumerate(['0–499','500–999','…','4500–4999','0–499']):
    x=.99+i*2.44
    rect(s,x,5.94,2.09,.48,WHITE,radius=True)
    text(s,x+.05,6.04,1.99,.28,val,17,TEAL,True,font=MONO,align=PP_ALIGN.CENTER)
    if i<4: arrow(s,x+2.14,6.09,.2,.15)
text(s,.99,6.53,11.4,.22,'各バッチは S で進行。最後の矢印では、平文番号と乱数シードをリセット。',15,MUTED)
note(s,'受信文字はfw_serial_readの戻り値としてr24に入っています。cpiでSと比較し、等しいとゼロフラグが立ち、breqで次の2:へ進みます。等しくなければretでmeasurement_loopから戻ります。小文字sや改行は開始コマンドになりません。2fはforward、1bのbはbackwardです。バッチ間ではxorshift_stateとcurrent_trace_indexを保持し、5000平文が終わった後の次のSで両方を初期化します。')

# 15
s=base('STEP 5 / Arduino 標準機能へつなぐ','計算と進行管理は .S、機器操作は Arduino へ','全体を動かすには、同じフォルダーにある .S と .ino の両方が必要です。')
label(s,.88,2.85,3.52,'measurement.S','AES・乱数・平文生成\n繰り返し・再開・ログ順序',h=1.8,fill=PALE)
arrow(s,4.56,3.59,.47,.28)
label(s,5.18,2.85,3.12,'.ino の接続関数','アセンブリから\n標準機能を呼ぶ窓口',h=1.8)
arrow(s,8.49,3.59,.47,.28)
label(s,9.13,2.85,3.35,'Arduino Core','Serial / digitalWrite\npinMode / delay',h=1.8)
code(s,.95,5.18,5.0,['call fw_trigger_high'],23,title='アセンブリ側')
code(s,6.45,5.18,5.93,['void fw_trigger_high() {','  digitalWrite(5, HIGH);','}'],19,title='接続用 .ino 側')
note(s,'機器操作の標準実装は維持しました。Serialの送受信バッファ、digitalWriteのピン処理、delayのタイマー処理を独自に書き直していません。callは関数を呼び出す命令で、retで戻ります。extern "C"はC++側の関数名をアセンブリから呼べるようにするための指定です。Arduinoの初期化、割り込み、起動処理も残っています。今回の成果物は接続用C++を伴うAVRアセンブリ版であり、すべてのランタイムを手書きした単独の.Sではありません。','Standard_AES_Measurement_ASM.ino:1–26; measurement.S:166–169')

# 16
s=base('STEP 5 / 機能の一致と測定波形','維持したのは、処理結果と測定の手順','命令列が変わるため、実行時間・トリガー幅・消費電力波形は変わります。')
steps=[('state へコピー',2.15),('D5 HIGH',2.05),('AES 計算',2.25),('D5 LOW',2.05),('10 ms 待機',2.25)]
x=.9
for head,w in steps:
    rect(s,x,2.86,w,.85,PALE if 'AES' in head else WHITE,radius=True)
    text(s,x+.07,3.09,w-.14,.34,head,20,TEAL,True,align=PP_ALIGN.CENTER)
    x+=w+.27
for y,name,end in [(4.56,'元コード',7.25),(5.58,'アセンブリ版',6.15)]:
    text(s,1.02,y-.2,2.65,.5,name,22)
    for coords in [(3.8,y,4.7,y),(4.7,y,4.7,y-.55),(4.7,y-.55,end,y-.55),(end,y-.55,end,y),(end,y,10.7,y)]:
        line(s,*coords,TEAL,3)
text(s,.98,6.45,11.4,.3,'模式図（実測値ではありません）。速度向上・実波形の一致は、今回の検証では主張しません。',17,MUTED)
note(s,'どちらもstateを平文からコピーし、トリガーを上げ、1ラウンドを実行し、トリガーを下げて10ms待ちます。それを同じ平文で10回繰り返し、ログの後にも10ms待ちます。図の幅の違いは模式的なもので、どちらが速いかを測定した結果ではありません。アセンブリ化するとコンパイラ任せのコードとは命令列が変わり、呼び出しのオーバーヘッドもあるため、時間と波形は同一にはなりません。割り込みは原本と同様に有効です。')

# 17
s=base('STEP 6 / 一致を確かめる','元コードとの比較と、別実装で結果を確認','以下は AVR シミュレータでの検証です。物理的な Uno と実波形は未検証です。','根拠：tests/run.py、tests/check.py、build/asm-test/stats.json、build/uno/stats.json')
for x,value,unit,body in [( .9,'5,500','平文','全5000 ＋ リセット後500'),(4.99,'55,000','パルス','各平文を10回測定'),(9.08,'60,500','待機要求','すべて delay(10)')]:
    rect(s,x,2.9,3.37,2.0,WHITE,radius=True)
    text(s,x+.17,3.12,3.03,.7,value,38,TEAL,True)
    text(s,x+.19,3.97,2.98,.35,unit,22,INK,True)
    text(s,x+.18,4.46,3.02,.26,body,14,MUTED)
text(s,1.0,5.3,11.35,1.22,'一致：ログ全文、各パルスの入力・出力、乱数列、再開・リセット\n追加確認：独立した Python 実装でも AES の結果を照合\n実機用バイナリ：最初の平文で10パルス、待機の最短間隔は約10.019 ms',20)
note(s,'比較用ビルドは長い待機だけを省略して、delay(10)の呼び出し回数と引数を記録しました。digitalWriteは実際のArduino実装を呼び、D5のパルス時点でSRAMの状態を採取しています。5000+500=5500平文、各10回で55000パルス、各平文10回の待機とログ後1回の待機で5500×11=60500回です。0に設定したテストや未実施の試験を成功に数えたものではありません。待機を置換しない実機用バイナリも同じシミュレータで最初の1平文を実行し、10パルスと実際のタイマー待機を確認しました。最短低レベル間隔160300サイクル÷16MHz=10.01875msです。実機上の実測ではありません。')

# 18
s=base('STEP 6 / 報告と読み進め方','「何を・どう・なぜ」で説明できるようにする','報告資料に使える説明例と、ソースを読む順番です。')
for y,head,body in [(2.7,'何を変えたか','AES・乱数生成・測定の進行管理を、AVR アセンブリへ移した。'),(3.82,'どう変えたか','C++ の処理を、読み込み・計算・保存・分岐の命令へ分解した。'),(4.94,'なぜその形か','Uno の8ビットCPUで実行する操作を、明示できる形にした。')]:
    rect(s,.91,y,11.53,.95,WHITE,radius=True)
    text(s,1.13,y+.17,2.23,.36,head,20,TEAL,True)
    text(s,3.55,y+.2,8.58,.35,body,20)
text(s,1.0,6.3,11.35,.36,'読む順番：execute_aes_std_asm → add_round_key_asm → 他の演算 → measurement_loop',17,MUTED)
note(s,'報告では単にアセンブリ化したと言うより、対象範囲、具体的な変更、理由、検証範囲の4点を述べると伝わります。演算結果と測定の進め方は維持、タイミングや消費電力波形は別途測定が必要という点も付け加えます。読み始めは全体の呼び出し順が短く書かれたexecute_aes_std_asm、次に具体例としてadd_round_key_asm、その後に各AES演算、最後にバッチを管理するmeasurement_loopの順がお勧めです。')

# 19
s=base('APPENDIX A / 命令の早見表','最初は、この12種類を見分ける','意味はこの資料のコードでの使い方に合わせています。')
entries=[('ldi','数値を直接入れる'),('ld / lds','SRAM の値を読む'),('st / sts','SRAM に値を保存する'),('lpm','フラッシュの値を読む'),('mov / movw','値をコピーする'),('eor','XOR を計算する'),('lsl / rol','左シフト・桁を引き継ぐ'),('cpi / cp','値を比較する'),('dec / inc','1 減らす・1 増やす'),('breq / brne','等しい・等しくないとき分岐'),('call / ret','関数を呼ぶ・戻る'),('push / pop','値を一時退避・取り出す')]
for i,(inst,desc) in enumerate(entries):
    col=i//6; row=i%6; x=.92+col*6.15; y=2.73+row*.59
    rect(s,x,y,5.54,.49,WHITE,radius=True)
    text(s,x+.13,y+.07,1.92,.3,inst,18,TEAL,True,font=MONO)
    text(s,x+2.15,y+.06,3.25,.32,desc,17)
note(s,'これは完全な命令リファレンスではありません。分岐は直前の命令が設定したフラグを見ます。brneは常にレジスタが0でないかを直接比較する命令ではなく、ゼロフラグが0のときに分岐します。movwは隣接した2レジスタのペアをコピーします。rolはキャリーを経由した左回転で、複数バイトをつなぐシフトに使っています。命令の細かい定義はMicrochipのAVR Instruction Set Manualで確認できます。')

# 20
s=base('APPENDIX B / 記号と呼び出しの約束','ドット・ラベル・退避も、役割を分けて読む','CPU が実行する命令と、プログラムを組み立てる指定が混在します。')
items=[('.equ / .byte / .section','定数・データ・配置先を指定する。CPU の演算命令ではない。'),('label: / 1b / 2f','処理の目印。b は後方、f は前方の同じ番号のラベルを指す。'),('lo8(...) / hi8(...)','16ビットのアドレスを、下位8ビット・上位8ビットに分ける。'),('push / pop','関数呼び出しで変わり得る値を退避し、逆順に取り戻す。'),('r1 = 0 / r24 など','Arduino とつなぐため、AVR GCC の呼び出し規約に従う。')]
for i,(head,desc) in enumerate(items):
    y=2.7+i*.74
    rect(s,.94,y,11.5,.62,WHITE,radius=True)
    text(s,1.1,y+.12,3.55,.34,head,18,TEAL,True,font=MONO)
    text(s,4.93,y+.13,7.2,.33,desc,17)
note(s,'GNU AVRアセンブラのディレクティブはソースを組み立てる指定です。.equで定数を定義し、.byteでデータを置き、.sectionでコードやデータの区分を指定します。push/popはスタックを使う実行命令です。今回の平文生成ループでは、Xポインタとカウンタをpushで保存してxorshiftを呼び、popで逆順に戻しています。r1のゼロ維持や、呼び出し側・呼び出される側が守るレジスタの約束に従うことでArduino C++と接続できます。今回の関数はcall-clobberedのr18〜r27、r30〜r31などを使用しています。')

# 21
s=base('APPENDIX C / ソースと根拠','元コードとアセンブリを、同じ処理で探す','行番号は今回作成したファイルのものです。')
rows=[('元コードの処理','measurement.S の対応先','行'),('execute_aes_std','execute_aes_std_asm','285–292'),('AddRoundKey_Std','add_round_key_asm','297–315'),('SubBytes_Std','sub_bytes_asm','319–332'),('ShiftRows_Std','shift_rows_asm','336–361'),('MixColumns_Std','mix_columns_asm / XTIME','364–413'),('xorshift32','xorshift32_asm','230–281'),('setup / loop','measurement_setup / loop','40–212')]
for i,(a,b,c) in enumerate(rows):
    y=2.61+i*.365
    rect(s,.95,y,11.45,.34,TEAL if i==0 else (WHITE if i%2 else PALE))
    col=WHITE if i==0 else INK
    text(s,1.08,y+.04,3.7,.25,a,14,col,font=FONT if i==0 else MONO)
    text(s,5.03,y+.04,5.13,.25,b,14,col,font=FONT if i==0 else MONO)
    text(s,10.65,y+.04,1.5,.25,c,14,col,font=FONT if i==0 else MONO)
text(s,1.0,5.84,11.4,.75,'資料：添付原本、measurement.S、接続用 .ino、tests/check.py、検証結果 JSON\n命令の定義：Microchip「AVR Instruction Set Manual」／GNU as「AVR Dependent Features」',15,MUTED)
note(s,'主な根拠はユーザー添付の原本と、この会話で作成したソースコード・テスト出力です。今回のアセンブリZIPはGitHubのcodex/assembly-zip-downloadブランチに保存されています。再現用の詳細はZIP内のREADME_ASM.mdを参照します。検証済みツールはAVR GCC 14.2.0、Arduino AVR Core 1.8.7、simavr 1.6です。参照先：https://github.com/toritenn27-netizen/Standard_AES_Measurement/tree/codex/assembly-zip-download 。命令の公式定義：https://www.microchip.com/en-us/product/ATmega328P のドキュメント、および https://sourceware.org/binutils/docs/as/AVR_002dDependent.html 。外部資料は説明の補助としての参照先で、今回の動作一致の根拠は実際のソースと検証記録です。')

prs.core_properties.title='Arduino Uno の AES をアセンブリで読む'
prs.core_properties.subject='初学者向け：コードの変更点、AVR命令の読み方、機能一致の検証'
prs.core_properties.author='Codex'
prs.core_properties.keywords='Arduino Uno, AVR, AES, アセンブリ, 初学者'
target=OUT/'AES_Assembly_Beginners.pptx'
prs.save(target)
(OUT/'speaker_notes.md').write_text('# 発表者ノート\n\n'+ '\n\n'.join(f'## {n:02d} {title}\n\n{body}' for n,title,body in notes),encoding='utf-8')
tiny=[item for item in checks if item[2]<13 and item[1] and not item[1].startswith(('Arduino Uno /','原本','根拠','measurement.S','対象：'))]
print('Slides:',len(prs.slides))
print('Saved:',target)
print('Small text audit:',tiny)
